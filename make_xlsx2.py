from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
wb = Workbook()
HF = Font(bold=True, color="FFFFFF"); HFill = PatternFill("solid", fgColor="2F4F4F")
def header(ws, row, labels):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab); c.font = HF; c.fill = HFill
        c.alignment = Alignment(wrap_text=True, vertical="center")
def widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[get_column_letter(i)].width = x
def wrap(ws, rng):
    for row in ws[rng]:
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")

# 1. V_n check (formulas)
ws = wb.active; ws.title = "V_n check"
ws["A1"] = "Hadamard V_n on de Sitter (H = 1): symbolic recursion vs numerical fit"; ws["A1"].font = Font(bold=True, size=12)
header(ws, 3, ["m²/H²","V0 predicted","V0 numeric","V0 diff","V1 predicted","V1 numeric","V1 diff","V2 predicted","V2 numeric","V2 diff"])
rows = [(0.3,-0.06375,-0.0122187500006,-0.85),(1,-0.125,-0.0312499999924,-0.5),(3,0.375,0.156249999899,0.5),(4,1.0,0.499999999611,1.0)]
for r,(m2,v1n,v2n,v0n) in enumerate(rows, start=4):
    ws.cell(r,1,m2); ws.cell(r,2,f"=(A{r}-2)/2"); ws.cell(r,3,v0n); ws.cell(r,4,f"=ABS(C{r}-B{r})")
    ws.cell(r,5,f"=A{r}*(A{r}-2)/8"); ws.cell(r,6,v1n); ws.cell(r,7,f"=ABS(F{r}-E{r})")
    ws.cell(r,8,f"=A{r}*(A{r}-2)*(A{r}+2)/96"); ws.cell(r,9,v2n); ws.cell(r,10,f"=ABS(I{r}-H{r})")
    for c in range(2,11): ws.cell(r,c).number_format = "0.000000000000"
ws["A9"] = "Predicted columns are formulas in m². Numeric values from check_Vn.py (120-digit least-squares fit)."
widths(ws, [10,14,14,12,14,14,12,14,14,12])

# 2. W0 check
ws2 = wb.create_sheet("W0 check")
ws2["A1"] = "W0 on de Sitter (H = 1): numeric vs closed form (constant fixed by antipodal regularity)"; ws2["A1"].font = Font(bold=True, size=12)
header(ws2, 3, ["m²/H²","h± type","ψ(h+)+ψ(h-)","W0 numeric","W0 closed form","diff"])
w0 = [(0.3,7.93808006776867,-9.192358738774892,7.938080067768674),(0.5,3.90611226700736,-5.113878283030708,3.906112267007357),
      (1,0.905336287209548,-1.938623390328882,0.9053362872095476),(2,-0.333333333333333,-0.1544313298030657,-0.3333333333333333),
      (3,-0.317617817441,0.5701468825406071,-0.3176178174414696),(4,0.133091642954088,1.005140827044301,0.1330916429540882)]
for r,(m2,num,S,closed) in enumerate(w0, start=4):
    ws2.cell(r,1,m2); ws2.cell(r,2,"real" if m2<=2.25 else "complex pair")
    ws2.cell(r,3,S); ws2.cell(r,4,num); ws2.cell(r,5,closed); ws2.cell(r,6,f"=ABS(E{r}-D{r})")
    for c in range(3,7): ws2.cell(r,c).number_format = "0.000000000000"
widths(ws2, [10,16,18,18,18,12])

# 3. Candidate tests
ws3 = wb.create_sheet("Candidate tests")
ws3["A1"] = "Kernel candidates: R2 (positivity) and R3 (coincidence) status"; ws3["A1"].font = Font(bold=True, size=12)
header(ws3, 3, ["Candidate","R2 positivity","R3 as originally written","Reworded R3 (coincidence local)","Key numbers","Status"])
cands = [
 ("H_ren (paper)","Hadamard state, not tested here","FAILS (W0 ≠ 0)","FAILS (non-local constant at coincidence)","W0 = 0.905 at m²=1","Original kernel"),
 ("H_ren − ⟨φ²⟩ (subtraction)","Mixed sign: +0.0067, +0.0082, −0.079, −0.15 (m² = 0.3, 1, 3, 4)","Would pass by construction","Not tested","Homogeneous-eq residual at μ→0: +0.0040, +0.0032, −0.0203, −0.0300","REJECTED"),
 ("G_S = (−∇²+m²)⁻¹","HOLDS: spectral weights > 0; spectral sum = closed form to ~1e-4","FAILS (1/σ and W0 ≠ 0)","Not local at coincidence as an inverse","∫G dV = 1/m² exactly","Positive; R3 fails"),
 ("K_t = e^{−tL}","HOLDS by construction","FAILS (K_t(0) ≠ 0 at finite t)","HOLDS: a1 = 2 − m², a2 = −1/15 + (2−m²)²/2","a2 numeric 0.43333 at m²=1 (error ∝ t)","Passes reworded R3"),
 ("Entropy-weighted L = −∇²+m²+κS","Trivial on constant background (mass shift)","n/a","n/a","S_Wald = S_BH constant","Untested (needs φ̇ ≠ 0)"),
]
for r,row in enumerate(cands, start=4):
    for c,v in enumerate(row,1): ws3.cell(r,c,v)
wrap(ws3, f"A4:F{3+len(cands)}")
widths(ws3, [26,38,26,34,38,22])

# 4. Summary
ws4 = wb.create_sheet("Summary")
ws4["A1"] = "IRFormal1: status summary (de Sitter, constant φ̄, mixing absent)"; ws4["A1"].font = Font(bold=True, size=12)
header(ws4, 3, ["Item","Result","Status"])
items = [
 ("Dictionary","f(φ)G^{μν}∂φ∂φ = Horndeski G₅ with G₅ = M_P²F(φ), F′ = f; paper's G₄ label incorrect","DERIVED"),
 ("Kinetic normalisation","Z = −6H²M_P²f̄; healthy requires f̄ < 0","DERIVED"),
 ("Mixing δφ–h","Zero at quadratic order on this background","DERIVED"),
 ("V₀, V₁, V₂","(m²−2)/2; m²(m²−2)/8; m²(m²−2)(m²+2)/96","DERIVED, CHECKED"),
 ("W₀","Closed form; constant fixed by antipodal regularity","CHECKED (regularity argument analytic)"),
 ("ℓ-dependence","W_ℓ = W₁ + V₀ ln ℓ²","DERIVED"),
 ("Wald entropy / horizon damping","S = S_BH on constant background; damping ≡ 1 (test trivial)","DERIVED; coefficient for G₅ DEFERRED"),
 ("R3 (original)","Fails: W₀ ≠ 0","CHECKED"),
 ("R3 (reworded: coincidence local)","Heat kernel passes; G_S does not","CHECKED"),
 ("Heat-kernel time t","Unfixed","OPEN"),
 ("Lorentzian R2","Not established for K_t or G_S","OPEN"),
 ("Time-dependent background, horizon test","Outside scope of IRFormal1","DEFERRED to paper two"),
]
for r,row in enumerate(items, start=4):
    for c,v in enumerate(row,1): ws4.cell(r,c,v)
wrap(ws4, f"A4:C{3+len(items)}")
widths(ws4, [34,70,34])

out = "IRFormal1_results.xlsx"; wb.save(out); print("saved", out)
